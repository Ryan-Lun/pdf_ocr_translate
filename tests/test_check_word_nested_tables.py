from __future__ import annotations

import docx

from scripts import check_word_nested_tables


def test_scan_docx_finds_nested_table_in_late_large_table_row(tmp_path):
    source_path = tmp_path / "late_nested.docx"
    document = docx.Document()
    table = document.add_table(rows=60, cols=3)
    for row_index, row in enumerate(table.rows, start=1):
        row.cells[0].text = f"row {row_index}"
    nested_table = table.rows[-1].cells[0].add_table(rows=1, cols=1)
    nested_table.cell(0, 0).paragraphs[0].text = "內層表格"
    document.save(source_path)

    findings = check_word_nested_tables.scan_docx(source_path)

    assert len(findings) == 1
    assert findings[0].location == "body"
    assert findings[0].top_table_index == 1
    assert findings[0].path == "body.table[1].r60c1.table[1]"
    assert findings[0].depth == 1
