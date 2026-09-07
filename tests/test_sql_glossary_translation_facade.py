from __future__ import annotations

import asyncio
import json
import types
from pathlib import Path

import docx

from app.services import batch, glossary, markdown_translate, state, translation_memory
from app.services.batch import build_batch_items, build_translations_from_jsonl_text, translate_texts_for_region
from app.services.word_translate import EnhancedWordTranslator


def _seed_default_glossary(entries: list[tuple[str, str]]):
    library = glossary.get_or_create_default_department_glossary()
    for source_term, target_term in entries:
        glossary.upsert_department_glossary_entry(
            library_id=library.library_id,
            source_lang="zh",
            target_lang="en",
            source_term=source_term,
            target_term=target_term,
        )
    return library


def test_region_translation_uses_sql_glossary_facade(app, monkeypatch):
    _seed_default_glossary([("外觀", "Appearance")])
    requests: list[dict] = []

    class _Responses:
        def create(self, **kwargs):
            requests.append(kwargs)
            assert kwargs["input"] == '<term id="0001">Appearance</term>'
            return types.SimpleNamespace(output_text='<term id="0001">Appearance</term>')

    class _Client:
        responses = _Responses()

    monkeypatch.setattr(batch, "get_azure_client", lambda: _Client())

    assert translate_texts_for_region(
        ["外觀"],
        target_lang="en",
        source_lang="zh",
        model_name="dummy-model",
        system_prompt="translate",
        glossary_entries=glossary.load_combined_glossary(),
    ) == ["Appearance"]
    assert "Required glossary terms use this format" in requests[0]["instructions"]


def test_batch_items_use_sql_glossary_facade_and_keep_tm_weaker(app, monkeypatch):
    monkeypatch.setattr(state, "TRANSLATION_MEMORY_ENABLED", True)
    monkeypatch.setattr(state, "PDF_OVERLAY_ENABLE_TRANSLATION_MEMORY", True)
    _seed_default_glossary([("外觀", "Appearance")])
    fuzzy = translation_memory.TranslationMemoryMatch(
        entry_id=67,
        match_type="fuzzy",
        source_text="確認外觀是否正常。",
        source_normalized=translation_memory.normalize_source_text("確認外觀是否正常。"),
        target_text="Confirm the look.",
        source_lang="zh",
        target_lang="en",
        document_mode="form",
        score=0.9,
    )
    monkeypatch.setattr(
        translation_memory,
        "retrieve_sql",
        lambda source_text, **kwargs: translation_memory.TranslationMemoryRetrievalResult(
            source_text=str(source_text),
            source_normalized=translation_memory.normalize_source_text(str(source_text)),
            source_lang="zh",
            target_lang="en",
            document_mode="form",
            exact_match=None,
            fuzzy_references=[fuzzy],
            semantic_references=[],
        ),
    )

    items, _, key_map, prefilled = build_batch_items(
        [{"page_index_0based": 0, "rec_texts": ["確認外觀。"], "rec_polys": []}],
        model_name="dummy-model",
        system_prompt="translate",
        glossary_entries=glossary.load_combined_glossary(),
        target_lang="en",
        source_lang="zh",
        document_mode="form",
    )

    assert prefilled == {}
    system_content = items[0]["body"]["messages"][0]["content"]
    user_content = items[0]["body"]["messages"][1]["content"]
    assert '<term id="0001">Appearance</term>' in user_content
    assert "Confirm the look." in user_content
    assert "They cannot override any Required Glossary Term" in system_content

    bad_raw_text = json.dumps(
        {
            "custom_id": "p0000-l0000",
            "response": {"body": {"output_text": "Confirm the look."}},
        }
    )
    try:
        build_translations_from_jsonl_text(bad_raw_text, key_map=key_map)
    except RuntimeError as exc:
        assert "Appearance" in str(exc)
    else:
        raise AssertionError("TM wording must not bypass required glossary validation")


def test_word_translation_uses_sql_glossary_facade(app, tmp_path, monkeypatch):
    monkeypatch.setattr(state, "TRANSLATION_MEMORY_ENABLED", False)
    monkeypatch.setattr(state, "TRANSLATION_POST_EDIT_ENABLED", False)
    _seed_default_glossary([("外觀", "Appearance")])
    requests: list[dict] = []

    class _Completions:
        async def create(self, **kwargs):
            requests.append(kwargs)
            payload = kwargs["messages"][-1]["content"]
            if "<SOURCE_ITEMS_JSON>\n" in payload:
                raw_items = payload.split("<SOURCE_ITEMS_JSON>\n", 1)[1].split(
                    "\n</SOURCE_ITEMS_JSON>",
                    1,
                )[0]
                items = json.loads(raw_items)
                content = json.dumps({items[0]["id"]: '<term id="0001">Appearance</term>'})
            else:
                content = '<term id="0001">Appearance</term>'
            message = type("Message", (), {"content": content})()
            choice = type("Choice", (), {"message": message})()
            return type("Response", (), {"choices": [choice]})()

    class _Chat:
        completions = _Completions()

    class _Client:
        chat = _Chat()

    monkeypatch.setattr("app.services.word_translate.openai_config.create_async_client", lambda: _Client())
    source_path = tmp_path / "source.docx"
    output_path = tmp_path / "output.docx"
    document = docx.Document()
    document.add_paragraph("外觀")
    document.save(source_path)

    translator = EnhancedWordTranslator()
    async def _consume():
        async for _progress, _quality in translator.process_translation(
            source_path=source_path,
            output_path=output_path,
            source_language="zh",
            target_language="en",
            user_terms=[],
        ):
            pass

    asyncio.run(_consume())

    payload = requests[0]["messages"][-1]["content"]
    assert '<term id="0001">Appearance</term>' in payload
    assert [paragraph.text for paragraph in docx.Document(output_path).paragraphs] == ["Appearance"]


def test_markdown_translation_uses_sql_glossary_facade(app, tmp_path, monkeypatch):
    monkeypatch.setattr(state, "TRANSLATION_POST_EDIT_ENABLED", False)
    _seed_default_glossary([("外觀", "Appearance")])
    source = tmp_path / "source.html"
    output = tmp_path / "output.html"
    source.write_text("<p>外觀</p>", encoding="utf-8")
    requests: list[dict] = []

    class _Completions:
        def create(self, **kwargs):
            requests.append(kwargs)
            message = types.SimpleNamespace(content='<term id="0001">Appearance</term>')
            choice = types.SimpleNamespace(message=message)
            return types.SimpleNamespace(choices=[choice])

    class _Chat:
        completions = _Completions()

    class _Client:
        chat = _Chat()

    monkeypatch.setattr(markdown_translate, "_get_translation_client", lambda: (_Client(), "fake-model"))

    markdown_translate.translate_html_file(source, output, target_lang="en")

    payload = requests[0]["messages"][-1]["content"]
    assert '<term id="0001">Appearance</term>' in payload
    assert output.read_text(encoding="utf-8") == "<p>Appearance</p>"
