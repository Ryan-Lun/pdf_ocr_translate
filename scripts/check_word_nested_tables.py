from __future__ import annotations

import argparse
import csv
import json
import sys
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any, Iterable

import docx


@dataclass(frozen=True)
class NestedTableFinding:
    file: str
    location: str
    top_table_index: int
    path: str
    depth: int
    row_index: int
    column_index: int
    cell_text_preview: str


def _iter_input_files(paths: Iterable[Path]) -> tuple[list[Path], list[Path]]:
    docx_files: list[Path] = []
    unsupported_doc_files: list[Path] = []
    for raw_path in paths:
        path = raw_path.resolve()
        if path.is_dir():
            candidates = sorted(item for item in path.rglob("*") if item.is_file())
        else:
            candidates = [path]
        for candidate in candidates:
            name = candidate.name
            suffix = candidate.suffix.lower()
            if name.startswith("~$"):
                continue
            if suffix == ".docx":
                docx_files.append(candidate)
            elif suffix == ".doc":
                unsupported_doc_files.append(candidate)
    return sorted(set(docx_files)), sorted(set(unsupported_doc_files))


def _preview_cell_text(cell: Any, *, limit: int = 120) -> str:
    text = " ".join(paragraph.text.strip() for paragraph in cell.paragraphs if paragraph.text.strip())
    text = " ".join(text.split())
    if len(text) <= limit:
        return text
    return text[: limit - 3] + "..."


def _scan_table(
    table: Any,
    *,
    file_path: Path,
    location: str,
    top_table_index: int,
    table_path: str,
    depth: int,
) -> list[NestedTableFinding]:
    findings: list[NestedTableFinding] = []
    seen_cells: set[Any] = set()
    for row_index, row in enumerate(table.rows, start=1):
        for column_index, cell in enumerate(row.cells, start=1):
            cell_id = cell._tc
            if cell_id in seen_cells:
                continue
            seen_cells.add(cell_id)
            for nested_index, nested_table in enumerate(cell.tables, start=1):
                nested_path = f"{table_path}.r{row_index}c{column_index}.table[{nested_index}]"
                findings.append(
                    NestedTableFinding(
                        file=str(file_path),
                        location=location,
                        top_table_index=top_table_index,
                        path=nested_path,
                        depth=depth + 1,
                        row_index=row_index,
                        column_index=column_index,
                        cell_text_preview=_preview_cell_text(cell),
                    )
                )
                findings.extend(
                    _scan_table(
                        nested_table,
                        file_path=file_path,
                        location=location,
                        top_table_index=top_table_index,
                        table_path=nested_path,
                        depth=depth + 1,
                    )
                )
    return findings


def _scan_header_footer_part(
    part: Any,
    *,
    file_path: Path,
    location: str,
) -> list[NestedTableFinding]:
    findings: list[NestedTableFinding] = []
    for top_table_index, table in enumerate(part.tables, start=1):
        findings.extend(
            _scan_table(
                table,
                file_path=file_path,
                location=location,
                top_table_index=top_table_index,
                table_path=f"{location}.table[{top_table_index}]",
                depth=0,
            )
        )
    return findings


def scan_docx(path: Path, *, include_header_footer: bool = False) -> list[NestedTableFinding]:
    document = docx.Document(path)
    findings: list[NestedTableFinding] = []
    for top_table_index, table in enumerate(document.tables, start=1):
        findings.extend(
            _scan_table(
                table,
                file_path=path,
                location="body",
                top_table_index=top_table_index,
                table_path=f"body.table[{top_table_index}]",
                depth=0,
            )
        )
    if include_header_footer:
        for section_index, section in enumerate(document.sections, start=1):
            parts = (
                ("header", section.header),
                ("first_page_header", section.first_page_header),
                ("even_page_header", section.even_page_header),
                ("footer", section.footer),
                ("first_page_footer", section.first_page_footer),
                ("even_page_footer", section.even_page_footer),
            )
            for part_name, part in parts:
                findings.extend(
                    _scan_header_footer_part(
                        part,
                        file_path=path,
                        location=f"section[{section_index}].{part_name}",
                    )
                )
    return findings


def _write_csv(path: Path, findings: list[NestedTableFinding]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(asdict(findings[0]).keys()) if findings else [
            "file",
            "location",
            "top_table_index",
            "path",
            "depth",
            "row_index",
            "column_index",
            "cell_text_preview",
        ])
        writer.writeheader()
        for finding in findings:
            writer.writerow(asdict(finding))


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Scan .docx files for nested Word tables."
    )
    parser.add_argument("paths", nargs="+", type=Path, help=".docx file or directory to scan.")
    parser.add_argument(
        "--include-header-footer",
        action="store_true",
        help="Also scan nested tables in headers and footers.",
    )
    parser.add_argument("--json", action="store_true", help="Print JSON instead of text lines.")
    parser.add_argument("--csv", type=Path, default=None, help="Write findings to a CSV file.")
    parser.add_argument(
        "--fail-on-nested",
        action="store_true",
        help="Exit with code 2 when nested tables are found.",
    )
    args = parser.parse_args(argv)

    docx_files, unsupported_doc_files = _iter_input_files(args.paths)
    findings: list[NestedTableFinding] = []
    errors: list[dict[str, str]] = []
    for file_path in docx_files:
        try:
            findings.extend(scan_docx(file_path, include_header_footer=args.include_header_footer))
        except Exception as exc:
            errors.append({"file": str(file_path), "error": f"{type(exc).__name__}: {exc}"})

    if args.csv:
        _write_csv(args.csv, findings)

    summary = {
        "scanned_docx": len(docx_files),
        "unsupported_doc": len(unsupported_doc_files),
        "files_with_nested_tables": len({finding.file for finding in findings}),
        "nested_tables": len(findings),
        "errors": len(errors),
    }

    if args.json:
        print(
            json.dumps(
                {
                    "summary": summary,
                    "findings": [asdict(finding) for finding in findings],
                    "unsupported_doc_files": [str(path) for path in unsupported_doc_files],
                    "errors": errors,
                },
                ensure_ascii=False,
                indent=2,
            )
        )
    else:
        for finding in findings:
            print(
                "nested_table "
                f"file={finding.file} "
                f"location={finding.location} "
                f"top_table={finding.top_table_index} "
                f"path={finding.path} "
                f"depth={finding.depth} "
                f"cell=r{finding.row_index}c{finding.column_index} "
                f"preview={finding.cell_text_preview!r}"
            )
        for path in unsupported_doc_files:
            print(f"unsupported_doc file={path} reason=python-docx-cannot-read-binary-doc", file=sys.stderr)
        for error in errors:
            print(f"scan_error file={error['file']} error={error['error']}", file=sys.stderr)
        print(
            "summary "
            f"scanned_docx={summary['scanned_docx']} "
            f"unsupported_doc={summary['unsupported_doc']} "
            f"files_with_nested_tables={summary['files_with_nested_tables']} "
            f"nested_tables={summary['nested_tables']} "
            f"errors={summary['errors']}"
        )
        if args.csv:
            print(f"csv={args.csv}")

    if errors:
        return 1
    if args.fail_on_nested and findings:
        return 2
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
